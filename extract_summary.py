import glob
import os
import openpyxl
import pandas as pd
import json

files = sorted(glob.glob("*.xlsx"))

results = []

for f in files:
    xl = pd.ExcelFile(f)
    print(f"\nProcessing {f}...")
    file_info = {
        "filename": f,
        "sheets": xl.sheet_names,
        "metadata": {},
        "columns": [],
        "sample_rows": [],
        "kpis": {},
        "feature_importance": [],
        "models_ranking": [],
        "watchlist_count": 0,
        "total_rows": 0
    }
    
    # 1. Check for standard structure like ML_13-ML_27
    main_sheet = None
    for s in xl.sheet_names:
        if "Enriched" in s or "Output" in s or "Lots" in s or "Orders" in s:
            main_sheet = s
            break
    if not main_sheet:
        main_sheet = xl.sheet_names[0]
        
    df_raw = xl.parse(main_sheet, header=None, nrows=12)
    
    # Check if metadata is in rows 0-4
    first_val = str(df_raw.iloc[0, 0]).strip() if pd.notna(df_raw.iloc[0, 0]) else ""
    if "Insight ID" in first_val or "Use Case" in first_val or "USE CASE" in first_val:
        for r in range(min(6, len(df_raw))):
            k = str(df_raw.iloc[r, 0]).strip() if pd.notna(df_raw.iloc[r, 0]) else ""
            v = str(df_raw.iloc[r, 1]).strip() if pd.notna(df_raw.iloc[r, 1]) else ""
            if k and v:
                file_info["metadata"][k] = v
        # Find where table header is
        table_header_row = None
        for r in range(4, min(12, len(df_raw))):
            non_nulls = df_raw.iloc[r].dropna().count()
            if non_nulls > 3:
                table_header_row = r
                break
        if table_header_row is not None:
            df_table = xl.parse(main_sheet, header=table_header_row, nrows=5)
            file_info["columns"] = [str(c) for c in df_table.columns if not str(c).startswith("Unnamed:")]
            file_info["sample_rows"] = df_table.head(3).to_dict(orient="records")
    else:
        # Standard table without top metadata header block
        df_table = xl.parse(main_sheet, nrows=5)
        file_info["columns"] = [str(c) for c in df_table.columns if not str(c).startswith("Unnamed:")]
        file_info["sample_rows"] = df_table.head(3).to_dict(orient="records")

    # Check for Executive Summary sheet if present
    for s in xl.sheet_names:
        if "Executive" in s or "KPI" in s:
            try:
                df_exec = xl.parse(s, header=None)
                file_info["executive_sheet_text"] = df_exec.dropna(how="all").head(15).values.tolist()
            except Exception as e:
                pass
        if "Feature" in s:
            try:
                # check feature importance
                df_feat = xl.parse(s)
                file_info["feature_importance_sample"] = df_feat.dropna(how="all").head(10).to_dict(orient="records")
            except Exception as e:
                pass
        if "Model" in s:
            try:
                df_mod = xl.parse(s)
                file_info["model_ranking_sample"] = df_mod.dropna(how="all").head(6).to_dict(orient="records")
            except Exception as e:
                pass
        if "Watchlist" in s or "Failure" in s or "Default" in s or "Delayed" in s:
            wb = openpyxl.load_workbook(f, read_only=True)
            if s in wb.sheetnames:
                file_info["watchlist_count"] = wb[s].max_row - 1
            wb.close()

    # Get total rows of main sheet
    try:
        wb = openpyxl.load_workbook(f, read_only=True)
        file_info["total_rows"] = wb[main_sheet].max_row
        wb.close()
    except:
        pass

    results.append(file_info)

with open("summary_metadata.json", "w", encoding="utf-8") as out:
    json.dump(results, out, indent=2, default=str)

print("Saved summary_metadata.json successfully!")
