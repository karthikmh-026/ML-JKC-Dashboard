import glob
import os
import openpyxl
import pandas as pd

files = sorted(glob.glob("*.xlsx"))
print(f"Total Excel files found: {len(files)}")

for f in files:
    try:
        excel = pd.ExcelFile(f)
        wb = openpyxl.load_workbook(f, read_only=True)
        print(f"\n==========================================")
        print(f"FILE: {f} ({os.path.getsize(f) / (1024*1024):.2f} MB)")
        print(f"Sheets: {wb.sheetnames}")
        for s in excel.sheet_names:
            ws = wb[s]
            row_count = ws.max_row
            df = excel.parse(s, nrows=2)
            cols = list(df.columns)
            print(f"  [Sheet: {s}] Rows: {row_count}, Cols ({len(cols)}): {cols[:7]}")
        wb.close()
    except Exception as e:
        print(f"Error {f}: {e}")
